"""TEXO XLSX KIT – bảng tính theo nhận diện TEXO (openpyxl).

Dùng:
    import sys; sys.path.insert(0, "<skill>/scripts"); import texo_xlsx_kit as X
    wb = X.new_wb()
    ws = wb.active
    r = X.setup_sheet(ws, "BẢNG THEO DÕI …", "Dòng phụ …", paper="A4", landscape=True)
    r = X.kpi_block(ws, r, [("12", "Hồ sơ"), ("3", "Phải xử lý", "high")])
    r = X.write_table(ws, r, ["#", "Nội dung", "Mức độ"], rows, widths=[6, 60, 16], sev_col=2)
    wb.save("out.xlsx")
Ô mức độ nhận "high" | "mid" | "ok" | "info" → icon ● + nhãn + nền màu.
"""
import os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ASSETS = os.path.join(os.path.dirname(__file__), "..", "assets")
LOGO = os.path.join(ASSETS, "logo_texo.png")
FONT = "Times New Roman"
C = {"indigo": "0C0150", "navy": "1F3864", "blue": "2E5496", "cyan": "01A0C7", "first": "EAF0F7",
     "zebra": "F7F9FC", "label": "D9E1F2", "grid": "BFBFBF", "grey": "595959"}
SEV = {"high": ("Phải xử lý", "C00000", "FDE4E1"), "mid": ("Nên xử lý", "C45911", "FFF1DE"),
       "ok": ("Đạt", "2E7D32", "E4F3E6"), "info": ("Lưu ý", "2E5496", "D9E1F2")}
THIN = Side(style="thin", color=C["grid"])
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def new_wb():
    return Workbook()


def setup_sheet(ws, title, subtitle="", paper="A4", landscape=True, ncols=8, logo=True,
                footer="Tài liệu nội bộ · Phòng Kỹ thuật TEXO"):
    """Logo + tiêu đề + cài đặt in (khổ, hướng, vừa chiều ngang, header/footer). Trả về dòng tiếp theo."""
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 20
    if logo and os.path.exists(LOGO):
        img = XLImage(LOGO)
        img.width, img.height = 121, 47  # giữ đúng tỷ lệ logo 900:349
        ws.add_image(img, "A1")
    start = 3 if logo else 1
    last = get_column_letter(ncols)
    ws.merge_cells(f"{get_column_letter(start)}1:{last}1")
    c = ws.cell(1, start, title)
    c.font = Font(name=FONT, size=15, bold=True, color=C["indigo"])
    c.alignment = Alignment(vertical="center")
    if subtitle:
        ws.merge_cells(f"{get_column_letter(start)}2:{last}2")
        c = ws.cell(2, start, subtitle)
        c.font = Font(name=FONT, size=10, italic=True, color=C["grey"])
    for col in range(1, ncols + 1):
        ws.cell(3, col).border = Border(bottom=Side(style="medium", color=C["blue"]))
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if paper == "A3" else ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddHeader.left.text = "CÔNG TY CP TEXO TƯ VẤN VÀ ĐẦU TƯ · PHÒNG KỸ THUẬT"
    ws.oddHeader.left.size = 8
    ws.oddHeader.right.text = title
    ws.oddHeader.right.size = 8
    ws.oddFooter.left.text = footer
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Trang &P / &N"
    ws.oddFooter.right.size = 8
    return 5


def kpi_block(ws, row, items, span=2):
    """Hàng thẻ KPI: items = [(giá trị, nhãn[, sev])]. Mỗi thẻ chiếm `span` cột."""
    for i, it in enumerate(items):
        val, lab = it[0], it[1]
        sev = it[2] if len(it) > 2 else None
        fg, bg = (SEV[sev][1], SEV[sev][2]) if sev else (C["indigo"], C["first"])
        c0 = 1 + i * (span + 0)
        ws.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + span - 1)
        ws.merge_cells(start_row=row + 1, start_column=c0, end_row=row + 1, end_column=c0 + span - 1)
        a = ws.cell(row, c0, val)
        a.font = Font(name=FONT, size=22, bold=True, color=fg)
        b = ws.cell(row + 1, c0, lab)
        b.font = Font(name=FONT, size=10, bold=True, color=C["navy"])
        for r in (row, row + 1):
            for cc in range(c0, c0 + span):
                ws.cell(r, cc).fill = PatternFill("solid", fgColor=bg)
                ws.cell(r, cc).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, c0).border = Border(top=Side(style="thick", color=fg))
    ws.row_dimensions[row].height = 34
    return row + 3


def write_table(ws, row, headers, rows, widths=None, sev_col=None, wrap=True, freeze=True, autofilter=True):
    """Bảng chuẩn: header nền 2E5496 chữ trắng, cột đầu EAF0F7, zebra, viền mảnh."""
    for j, h in enumerate(headers, 1):
        c = ws.cell(row, j, h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C["blue"])
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row + i, j)
            fill = C["first"] if j == 1 else (C["zebra"] if i % 2 == 0 else None)
            if sev_col is not None and j - 1 == sev_col and v in SEV:
                lab, fg, bg = SEV[v]
                c.value = "● " + lab
                c.font = Font(name=FONT, size=10, bold=True, color=fg)
                fill = bg
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.value = v
                c.font = Font(name=FONT, size=10)
                c.alignment = Alignment(vertical="center", wrap_text=wrap)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
            c.border = BOX
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    if freeze:
        ws.freeze_panes = ws.cell(row + 1, 2)
    if autofilter and rows:
        ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row + len(rows)}"
    ws.print_title_rows = f"{row}:{row}"
    return row + len(rows) + 2
