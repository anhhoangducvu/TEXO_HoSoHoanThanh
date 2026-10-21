# -*- coding: utf-8 -*-
"""Xuất kết quả kiểm tra ra Excel theo nhận diện TEXO (dùng texo_xlsx_kit)."""
import io
from datetime import datetime

from openpyxl.styles import Alignment, Font, PatternFill

from . import kb
from . import texo_xlsx_kit as X

SEVMAP = {"high": "high", "mid": "mid", "ok": "ok", "info": "info", "na": None}


def _sev(s):
    return SEVMAP.get(s) or "Không xét"


def xuat_xlsx(R, ten_file="", ai=None):
    wb = X.new_wb()
    meta, k = R["meta"], R["kpi"]
    sub = f"{meta.get('du_an', '')[:90]} · {R['loai_nt_label']} · {R['nghi_dinh']} · {datetime.now():%d/%m/%Y}"

    # ---------------- Tổng quan
    ws = wb.active
    ws.title = "Tổng quan"
    r = X.setup_sheet(ws, "KIỂM TRA HỒ SƠ CHẤT LƯỢNG DỰ ÁN", sub, paper="A4", landscape=True, ncols=8)
    r = X.kpi_block(ws, r, [(f"{k['muc_B_co']}/{k['muc_B']}", "Mục bắt buộc có"),
                            (f"{k['ty_le_co_ma']}%", "Dòng có mã VPĐT"),
                            (str(k["high"]), "Phải xử lý", "high"),
                            (str(k["mid"]), "Nên xử lý", "mid")], span=2)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    c = ws.cell(r, 1, "KẾT LUẬN: " + R["ket_luan"])
    lab, fg, bg = X.SEV.get(R["ket_luan_sev"], X.SEV["info"])
    c.font = Font(name=X.FONT, size=12, bold=True, color=fg)
    c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 36
    r += 2
    lab = [("tieu_de", "Tên tài liệu"), ("du_an", "Dự án"), ("hang_muc", "Hạng mục"), ("goi_thau", "Gói thầu"),
           ("dia_diem", "Địa điểm"), ("chu_dau_tu", "Chủ đầu tư"), ("tvgs", "Tư vấn giám sát"), ("nha_thau", "Nhà thầu thi công"),
           ("hop_dong", "Hợp đồng")]
    rows = [(l, meta.get(key, "")) for key, l in lab if meta.get(key)] + [("File kiểm tra", ten_file),
                                                                          ("Loại nghiệm thu", R["loai_nt_label"])]
    for a, b in rows:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.cell(r, 1, a).font = Font(name=X.FONT, size=10, bold=True, color=X.C["navy"])
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
        cc = ws.cell(r, 4, b)
        cc.font = Font(name=X.FONT, size=10)
        cc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 15 * max(1, -(-len(str(b or "")) // 150))
        r += 1
    r += 1
    frows = [[f["stt"], f["muc_do"], f["ma_loi"], f["tieu_de"], f["chi_tiet"], f.get("de_xuat", ""), f.get("muc_pl") or "",
              ", ".join(map(str, f["dong"][:30]))] for f in R["findings"]]
    X.write_table(ws, r, ["#", "Mức độ", "Mã lỗi", "Phát hiện", "Chi tiết", "Đề xuất", "Mục PL VII", "Dòng liên quan"],
                  frows, widths=[6, 16, 9, 48, 60, 45, 14, 22], sev_col=1, freeze=False)

    # ---------------- Đối chiếu PL VII
    ws = wb.create_sheet("Đối chiếu PL VII")
    r = X.setup_sheet(ws, "ĐỐI CHIẾU PHỤ LỤC VII – NĐ 207/2026", sub, ncols=10)
    rows = []
    for it in R["items"]:
        rows.append([it["code"], kb.NHOM[it["nhom"]], it["ten"], it["nguyen_van"], it["hinh_thuc_label"], it["yeu_cau"],
                     it.get("chu_the", ""), it["so_dong"], ", ".join(it["ma"]), _sev(it["trang_thai"]), it["ket_luan"]])
    X.write_table(ws, r, ["Mục", "Nhóm", "Thành phần hồ sơ", "Yêu cầu PL VII (rút gọn)", "Hình thức", "YC", "Chủ thể lưu",
                          "Số dòng", "Mã VPĐT", "Đánh giá", "Kết luận"], rows,
                  widths=[8, 26, 34, 52, 26, 5, 14, 8, 26, 14, 40], sev_col=9)

    # ---------------- Danh mục đã đọc
    ws = wb.create_sheet("Danh mục đã đọc")
    r = X.setup_sheet(ws, "DANH MỤC ĐÃ ĐỌC VÀ PHÂN LOẠI", sub, ncols=12)
    src = {"tu_khoa": "Từ khóa", "ke_thua": "Kế thừa nhóm", "ai": "AI"}
    rows = [[x["id"], x["cap"], x["stt"], ("    " * max(0, x["cap"] - 1)) + x["noi_dung"], x["so_vb"], x["ngay_raw"],
             x["don_vi"], ", ".join(x["ma"]), ", ".join(x["ma_ke_thua"]), x.get("muc") or "",
             src.get(x.get("muc_nguon"), ""), x["tinh_trang"], x["nguon"]] for x in R["rows"]]
    X.write_table(ws, r, ["ID", "Cấp", "STT", "Nội dung hồ sơ", "Số văn bản", "Ngày", "Đơn vị ban hành", "Mã VPĐT",
                          "Mã kế thừa", "Mục PL VII", "Nguồn phân loại", "Tình trạng", "Nguồn"], rows,
                  widths=[6, 5, 7, 70, 24, 12, 26, 18, 18, 10, 14, 12, 12])

    # ---------------- Mã VPĐT
    ws = wb.create_sheet("Mã VPĐT")
    r = X.setup_sheet(ws, "BẢN ĐỒ MÃ HỒ SƠ VPĐT", sub, ncols=7)
    rows = [[ma, len(e["dong"]), len(e["dong_ke_thua"]), ", ".join(sorted(c for c in e["muc"] if c)),
             "mid" if e.get("gom") else "ok", "; ".join(dict.fromkeys(e["noi_dung"]))[:300]]
            for ma, e in R["ma_map"].items()]
    X.write_table(ws, r, ["Mã HS", "Dòng gắn trực tiếp", "Dòng kế thừa", "Mục PL VII", "Đánh giá", "Nội dung"], rows,
                  widths=[18, 12, 12, 30, 16, 80], sev_col=4)
    dc = R.get("doi_chieu", {})
    if dc.get("bao_cao"):
        r2 = ws.max_row + 3
        ws.cell(r2, 1, "ĐỐI CHIẾU VỚI BÁO CÁO ĐĂNG TẢI").font = Font(name=X.FONT, size=12, bold=True, color=X.C["indigo"])
        rows = [[x, nd, m or "", "Có trong báo cáo, thiếu trong danh mục"] for x, nd, m in dc["bao_cao"]["chi_bc"]]
        rows += [[x, "", "", "Có trong danh mục, không có trong báo cáo"] for x in dc["bao_cao"]["chi_dm"]]
        X.write_table(ws, r2 + 1, ["Mã", "Nội dung (báo cáo)", "Mục PL VII", "Tình trạng"], rows, freeze=False, autofilter=False)

    if R.get("cau_hoi_muc"):
        ws = wb.create_sheet("Câu hỏi kiểm tra")
        r = X.setup_sheet(ws, "CÂU HỎI KIỂM TRA THEO TỪNG MỤC", sub, ncols=4)
        rows = []
        for q in R["cau_hoi_muc"]:
            for x in q["cau_hoi"]:
                rows.append([q["code"], q["ten"], x, q["trang_thai"] if q["code"] not in ("I", "III") else "info", ""])
        X.write_table(ws, r, ["Mục", "Thành phần", "Câu hỏi kiểm tra", "Hiện trạng", "Trả lời của đoàn"], rows,
                      widths=[8, 34, 80, 14, 50], sev_col=3)

    if ai and ai.get("ok"):
        ws = wb.create_sheet("Ý kiến AI")
        r = X.setup_sheet(ws, "Ý KIẾN BỔ SUNG CỦA TRỢ LÝ AI", sub, ncols=4)
        d = ai["data"]
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(r, 1, d.get("nhan_xet_tong_quan", ""))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=X.FONT, size=11)
        ws.row_dimensions[r].height = 120
        rows = [[i, x.get("muc_do", "info") if x.get("muc_do") in X.SEV else "info", x.get("noi_dung", ""), x.get("de_xuat", "")]
                for i, x in enumerate(d.get("rui_ro_bo_sung") or [], 1)]
        r = X.write_table(ws, r + 2, ["#", "Mức độ", "Rủi ro", "Đề xuất"], rows, widths=[6, 16, 70, 60], sev_col=1, freeze=False)
        rows = [[i, q] for i, q in enumerate(d.get("cau_hoi_cho_doan") or [], 1)]
        X.write_table(ws, r, ["#", "Câu hỏi cho đoàn"], rows, freeze=False, autofilter=False)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
