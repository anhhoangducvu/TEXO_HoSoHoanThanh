# -*- coding: utf-8 -*-
"""MẪU DANH MỤC THAM KHẢO (Excel) – không bắt buộc. Đoàn có thể dùng nguyên bộ, dùng từng sheet,
hoặc giữ mẫu riêng của mình; công cụ đọc theo nội dung, không theo mẫu.
Nếu dùng mẫu này: cột "Mục PL VII" giúp công cụ xếp đúng mục ngay, không phải đoán."""
import io

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import kb
from . import texo_xlsx_kit as X

COLS = ["STT", "Mục PL VII", "Nội dung hồ sơ", "Số hiệu / Ký hiệu", "Ngày ban hành", "Đơn vị ban hành / lập",
        "Đợt", "Tình trạng", "Mã TEXO-E", "Ghi chú"]
WID = [6, 10, 62, 24, 13, 26, 7, 13, 16, 30]

SHEETS = [
    ("1. Phap ly", "DANH MỤC HỒ SƠ PHÁP LÝ (PL VII – phần I)", [
        ("VD", "I.2", "Ví dụ: Quyết định phê duyệt dự án đầu tư xây dựng …", "123/QĐ-…", "15/03/2025", "Chủ đầu tư", "", "Bản sao", "HS.2025.1234", ""),
        ("VD", "I.7", "Ví dụ: Giấy phép xây dựng số … (hoặc: Miễn GPXD theo điểm … khoản 2 Điều 89 Luật XD)", "45/GPXD", "20/04/2025", "Sở Xây dựng", "", "Bản sao", "HS.2025.1235", ""),
        ("VD", "I.8", "Ví dụ: Hợp đồng thi công xây dựng gói thầu … (ghi rõ HĐ thầu chính / thầu phụ)", "15/2025/HĐXD/…", "25/03/2025", "CĐT – Nhà thầu", "", "Bản gốc", "HS.2025.1236", ""),
        ("VD", "I.9", "Ví dụ: Hồ sơ năng lực, sơ đồ tổ chức, QĐ bổ nhiệm chỉ huy trưởng, danh sách máy móc", "", "", "Nhà thầu", "", "Bản sao", "HS.2025.1237", ""),
    ], ["I.1", "I.2", "I.3", "I.4", "I.5", "I.6", "I.7", "I.8", "I.9", "I.10"]),
    ("2. Thiet ke", "DANH MỤC HỒ SƠ KHẢO SÁT, THIẾT KẾ – DANH MỤC BẢN VẼ (PL VII – phần II)", [
        ("VD", "II.3", "Ví dụ: QĐ phê duyệt thiết kế bản vẽ thi công hạng mục …", "88/QĐ-…", "10/05/2025", "Chủ đầu tư", "", "Bản sao", "HS.2025.1300", ""),
        ("VD", "II.3", "Ví dụ: KC-01 – Mặt bằng móng", "KC-01", "10/05/2025", "Đơn vị thiết kế", "", "Bản gốc", "HS.2025.1301", "Thuộc QĐ 88/QĐ-…"),
    ], ["II.1", "II.2", "II.3", "II.4", "II.5"]),
    ("3. Thay doi thiet ke", "DANH MỤC THAY ĐỔI THIẾT KẾ (PL VII – III.1)", [
        ("VD", "III.1", "Ví dụ: Điều chỉnh cao độ đáy móng trục A–C (kèm văn bản chấp thuận số …)", "12/CV-…", "02/07/2025", "Chủ đầu tư", "", "Bản gốc", "HS.2025.1400",
         "Không có thay đổi thì ghi một dòng: 'Không có thay đổi thiết kế'"),
    ], ["III.1"]),
    ("4. NT vat lieu", "DANH MỤC NGHIỆM THU VẬT LIỆU ĐẦU VÀO (PL VII – III.4)", [
        ("VD", "III.4", "Ví dụ: Nghiệm thu thép D10–D25 lô 01 (kèm CO/CQ, chứng nhận hợp quy)", "01/NTVL", "05/06/2025", "TVGS – Nhà thầu", "1", "Bản gốc", "HS.2025.1500", ""),
    ], ["III.4"]),
    ("5. Thi nghiem", "DANH MỤC KẾT QUẢ THÍ NGHIỆM, QUAN TRẮC, ĐO ĐẠC (PL VII – III.5, III.7)", [
        ("VD", "III.5", "Ví dụ: Kết quả nén mẫu bê tông móng M1–M5 (so với mác thiết kế)", "PTN-0123", "28/06/2025", "Phòng thí nghiệm LAS-XD …", "1", "Bản gốc", "HS.2025.1600", ""),
    ], ["III.5", "III.7"]),
    ("6. NT cong viec", "DANH MỤC BIÊN BẢN NGHIỆM THU CÔNG VIỆC, GIAI ĐOẠN (PL VII – III.6)", [
        ("VD", "III.6", "Ví dụ: Nghiệm thu cốt thép móng M1–M5, trục 1–5", "05/NTCV", "12/06/2025", "TVGS – Nhà thầu", "1", "Bản gốc", "HS.2025.1700", ""),
    ], ["III.6"]),
    ("7. Hoan cong", "DANH MỤC BẢN VẼ HOÀN CÔNG (PL VII – III.2)", [
        ("VD", "III.2", "Ví dụ: HC-KC-01 – Mặt bằng hoàn công móng", "HC-KC-01", "30/06/2025", "Nhà thầu (TVGS xác nhận)", "1", "Bản gốc", "HS.2025.1800", ""),
    ], ["III.2"]),
    ("8. Thiet bi", "HỒ SƠ THIẾT BỊ LẮP ĐẶT, CHẠY THỬ (PL VII – III.8)", [
        ("VD", "III.8", "Ví dụ: Biên bản nghiệm thu lắp đặt thiết bị trạm bơm", "01/NTTB", "", "", "", "Bản gốc", "", ""),
        ("VD", "III.8", "Ví dụ: Biên bản chạy thử đơn động / chạy thử liên động", "02/CTĐĐ", "", "", "", "Bản gốc", "", ""),
    ], ["III.8"]),
    ("9. Nghiem thu - khac", "VẬN HÀNH, PCCC, TỒN TẠI, NGHIỆM THU HOÀN THÀNH, BÁO CÁO (PL VII – III.9 … III.16)", [
        ("VD", "III.10b", "Ví dụ: Biên bản nghiệm thu nội bộ hệ thống PCCC / văn bản của cơ quan PCCC", "", "", "", "", "Bản gốc", "", ""),
        ("VD", "III.13", "Ví dụ: Biên bản nghiệm thu hoàn thành hạng mục … (có thời điểm bảo hành, đánh giá PCCC)", "01/BBNT-HT", "", "", "", "Bản gốc", "", ""),
        ("VD", "T.1", "Ví dụ: Báo cáo hoàn thành công tác giám sát (Phụ lục IVb NĐ 207)", "", "", "TVGS TEXO", "", "Bản gốc", "", ""),
    ], ["III.9", "III.10a", "III.10b", "III.10c", "III.10d", "III.10đ", "III.10e", "III.10g", "III.11", "III.12",
        "III.13", "III.14", "III.15", "III.16", "T.1", "T.2", "T.3"]),
]


def _guide(ws):
    r = X.setup_sheet(ws, "MẪU DANH MỤC HỒ SƠ CHẤT LƯỢNG – THAM KHẢO", "Căn cứ Phụ lục VII NĐ 207/2026/NĐ-CP · Phòng Kỹ thuật TEXO",
                      paper="A4", landscape=False, ncols=6)
    lines = [
        ("Mẫu này để tham khảo, không bắt buộc.", True),
        ("Phụ lục VII chỉ yêu cầu có danh mục (bản vẽ, hoàn công, nghiệm thu…), không quy định mẫu. Đoàn được dùng mẫu riêng – công cụ "
         "kiểm tra đọc theo nội dung từng dòng, không theo mẫu.", False),
        ("Cách dùng linh hoạt", True),
        ("• Có thể lập một quyển chung, hoặc tách thành nhiều file rời theo loại và theo đợt (VD: DM_NT_vat_lieu_dot_2.xlsx, "
         "DM_hoan_cong_dot_1.xlsx) rồi nạp cùng lúc vào app 'Kiểm tra Hồ sơ chất lượng dự án'.", False),
        ("• Đặt tên file hoặc tiêu đề có loại danh mục và số đợt để app nhận diện: pháp lý, thiết kế, thay đổi thiết kế, vật liệu, "
         "thí nghiệm, nghiệm thu công việc, hoàn công, thiết bị.", False),
        ("Cột tối thiểu nên có", True),
        ("Nội dung hồ sơ · Số hiệu / Ký hiệu · Ngày ban hành · Đơn vị ban hành · Tình trạng (bản gốc / bản sao) · Mã TEXO-E. "
         "Cột 'Mục PL VII' (I.1 … III.16) là tùy chọn – có thì app xếp đúng mục ngay.", False),
        ("Lưu ý khi lập", True),
        ("• Mục nào không phát sinh (thay đổi thiết kế, sự cố, thiết bị…) ghi một dòng 'Không có …' để chứng minh đã xem xét.", False),
        ("• Mỗi thành phần hồ sơ nên có mã HS riêng trên VPĐT; chưa đăng tải thì để trống và ghi chú 'chưa đăng tải'.", False),
        ("• Dòng có STT = VD là dòng ví dụ – app tự bỏ qua; có thể xóa đi.", False),
        ("• Hồ sơ thanh toán, quyết toán lập thành bộ riêng, không đưa vào danh mục hồ sơ hoàn thành.", False),
    ]
    ws.column_dimensions["A"].width = 110
    for txt, b in lines:
        c = ws.cell(r, 1, txt)
        c.font = Font(name=X.FONT, size=11 if not b else 12, bold=b, color=X.C["indigo"] if b else "1F1F1F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 18 if b else max(18, 15 * (len(txt) // 105 + 1))
        r += 1 if not b else 1
    r += 1
    ws.cell(r, 1, "Danh sách mục Phụ lục VII").font = Font(name=X.FONT, size=12, bold=True, color=X.C["indigo"])
    r += 1
    for m in kb.MUC:
        if m["code"].startswith("T.9"):
            continue
        ws.cell(r, 1, f"{m['code']}  ·  {m['ten']}  –  {m['nguyen_van']}").font = Font(name=X.FONT, size=10)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True)
        r += 1


def _tong(ws):
    r = X.setup_sheet(ws, "DANH MỤC TỔNG – QUYỂN HỒ SƠ HOÀN THÀNH", "Điền thông tin chung bên dưới; mỗi mục PL VII ghi tên hồ sơ thực tế hoặc 'Không có'",
                      ncols=10)
    for lab in ("Dự án:", "Hạng mục / gói thầu:", "Địa điểm:", "Chủ đầu tư:", "Tư vấn giám sát:", "Nhà thầu thi công:", "Hợp đồng số:"):
        ws.cell(r, 1, lab).font = Font(name=X.FONT, size=10, bold=True, color=X.C["navy"])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
        ws.cell(r, 3).fill = PatternFill("solid", fgColor="FFF9E6")
        r += 1
    r += 1
    rows = []
    for g in ("I", "II", "III", "T"):
        rows.append([g, "", kb.NHOM[g] if g != "T" else "HỒ SƠ BÁO CÁO, NHẬT KÝ", "", "", "", "", "", "", ""])
        for m in kb.MUC:
            if m["code"].split(".")[0] != g or m["code"] == "T.9":
                continue
            rows.append(["", m["code"], "", "", "", "", "", "", "", m["ten"]])
    hdr = ["STT", "Mục PL VII", "Nội dung hồ sơ", "Số hiệu / Ký hiệu", "Ngày ban hành", "Đơn vị ban hành / lập",
           "Quyển số / Tệp", "Tình trạng", "Mã TEXO-E", "Thành phần theo PL VII (tham chiếu)"]
    X.write_table(ws, r, hdr, rows, widths=[6, 10, 50, 22, 13, 24, 14, 12, 16, 46])
    for i, row in enumerate(rows, 1):
        if row[0]:
            for j in range(1, 11):
                c = ws.cell(r + i, j)
                c.font = Font(name=X.FONT, size=10, bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=X.C["navy"])
        else:
            ws.cell(r + i, 10).font = Font(name=X.FONT, size=9, italic=True, color=X.C["grey"])


def tao_mau_xlsx():
    wb = X.new_wb()
    ws = wb.active
    ws.title = "Huong dan"
    _guide(ws)
    _tong(wb.create_sheet("0. Danh muc tong"))
    codes = [m["code"] for m in kb.MUC if m["code"] != "T.9"]
    ds = wb.create_sheet("_ds")
    for i, c in enumerate(codes, 1):
        ds.cell(i, 1, c)
    ds.sheet_state = "hidden"
    for name, title, vd, goi_y in SHEETS:
        ws = wb.create_sheet(name)
        r = X.setup_sheet(ws, title, "Tham khảo – không bắt buộc theo mẫu. Có thể lưu riêng sheet này thành file theo từng đợt.", ncols=10)
        rows = [list(x) for x in vd]
        for i in range(1, 16):
            rows.append([str(i), goi_y[0] if len(goi_y) == 1 else "", "", "", "", "", "", "", "", ""])
        X.write_table(ws, r, COLS, rows, widths=WID)
        for i, row in enumerate(rows, 1):
            if row[0] == "VD":
                for j in range(1, 11):
                    ws.cell(r + i, j).font = Font(name=X.FONT, size=10, italic=True, color="7F7F7F")
        dv = DataValidation(type="list", formula1=f"=_ds!$A$1:$A${len(codes)}", allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"Bản gốc,Bản sao,Bản điện tử,Chưa có"', allow_blank=True)
        ws.add_data_validation(dv)
        ws.add_data_validation(dv2)
        dv.add(f"B{r + 1}:B{r + 200}")
        dv2.add(f"H{r + 1}:H{r + 200}")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
