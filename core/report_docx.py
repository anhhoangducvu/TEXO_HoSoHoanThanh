# -*- coding: utf-8 -*-
"""Xuất PHIẾU KIỂM TRA DANH MỤC HỒ SƠ HOÀN THÀNH ra Word A4 theo nhận diện TEXO (python-docx)."""
import io
import os
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips, Cm

from . import kb
from .checker import SEV_ICON, SEV_LABEL

COMPANY = "CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ"
DEPARTMENT = "PHÒNG KỸ THUẬT"
SIGNER_TITLE = "TRƯỞNG PHÒNG KỸ THUẬT"
SIGNER_NAME = "Hoàng Đức Vũ"
CITY = "Hà Nội"
LOGO = os.path.join(os.path.dirname(__file__), "..", "assets", "logo_texo.png")

FONT = "Times New Roman"
EMOJI_FONT = "Segoe UI Emoji"
CW = 9638  # bề rộng nội dung A4 dọc, lề 2,0 cm
CW_L = 14570  # A4 ngang
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = "2E5496"
GREY = RGBColor(0x59, 0x59, 0x59)
FILL = {"high": "FDE4E1", "mid": "FFF1DE", "ok": "E4F3E6", "info": "D9E1F2", "na": "F2F2F2"}
FG = {"high": RGBColor(0xC0, 0, 0), "mid": RGBColor(0xC4, 0x59, 0x11), "ok": RGBColor(0x2E, 0x7D, 0x32),
      "info": RGBColor(0x2E, 0x54, 0x96), "na": GREY}


# ----------------------------------------------------------------------------- helpers
def _font(run, size=12, bold=False, color=None, it=False, emoji=False):
    run.font.name = EMOJI_FONT if emoji else FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = it
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), EMOJI_FONT if emoji else FONT)


def _runs(p, text, size=12, bold=False, color=None, it=False):
    """Tách emoji ra run riêng (font Segoe UI Emoji) để không vỡ hiển thị."""
    buf, is_e = "", None
    for ch in str(text):
        e = ord(ch) > 0x2500 and ch not in "–—“”…·•→≤≥"
        if is_e is None or e == is_e:
            buf += ch
        else:
            _font(p.add_run(buf), size, bold, color, it, emoji=is_e)
            buf = ch
        is_e = e
    if buf:
        _font(p.add_run(buf), size, bold, color, it, emoji=bool(is_e))


def para(container, text="", size=12, bold=False, color=None, it=False, align=None, sa=4, sb=0):
    p = container.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    if text:
        _runs(p, text, size, bold, color, it)
    return p


def heading(doc, text, size=13):
    p = para(doc, text, size=size, bold=True, color=NAVY, sa=4, sb=8)
    p.paragraph_format.keep_with_next = True
    return p


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def _cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def _borders(table, color="BFBFBF", sz=4, none=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}")
        x.set(qn("w:val"), "none" if none else "single")
        if not none:
            x.set(qn("w:sz"), str(sz))
            x.set(qn("w:color"), color)
        b.append(x)
    tblPr.append(b)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def _set_widths(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if j < len(widths):
            gc.set(qn("w:w"), str(widths[j]))
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Twips(w)


def table(doc, headers, rows, widths, size=10, sev_cols=None, zebra=True):
    """Bảng chuẩn: header nền xanh chữ trắng, viền mảnh, lặp tiêu đề khi sang trang."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(t)
    _cell_margins(t)
    hr = t.rows[0]
    _repeat_header(hr)
    for j, h in enumerate(headers):
        c = hr.cells[j]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(p, h, size, True, RGBColor(0xFF, 0xFF, 0xFF))
        _shade(c, BLUE)
    for i, r in enumerate(rows):
        row = t.add_row()
        _no_split(row)
        sev = r[-1] if isinstance(r, (list, tuple)) and r and isinstance(r[-1], dict) else None
        vals = r[:-1] if sev else r
        for j, v in enumerate(vals):
            c = row.cells[j]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            txt = "" if v is None else str(v)
            lines = txt.split("\n")
            for k, line in enumerate(lines):
                if k:
                    p = c.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                _runs(p, line, size)
            if sev and j in sev.get("cols", []):
                _shade(c, FILL.get(sev["sev"], "FFFFFF"))
                for pp in c.paragraphs:
                    for rr in pp.runs:
                        rr.font.color.rgb = FG.get(sev["sev"], GREY)
                        rr.font.bold = True
            elif zebra and i % 2 == 1:
                _shade(c, "F7F9FC")
    _set_widths(t, widths)
    para(doc, "", sa=2)
    return t


def _pagefield(p):
    r = p.add_run()
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = "PAGE"
    z = OxmlElement("w:fldChar"); z.set(qn("w:fldCharType"), "end")
    r._r.append(a); r._r.append(i); r._r.append(z)
    _font(r, 9, color=GREY)


def _pborder(p, edge="bottom", color=BLUE, sz=6, sp=4):
    pp = p._p.get_or_add_pPr()
    b = pp.find(qn("w:pBdr"))
    if b is None:
        b = OxmlElement("w:pBdr")
        pp.append(b)
    e = OxmlElement(f"w:{edge}")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", str(sp)), ("w:color", color)):
        e.set(qn(k), v)
    b.append(e)


def _header_footer(sec, cw, title):
    sec.header_distance = sec.footer_distance = Twips(500)
    hdr = sec.header
    hdr.is_linked_to_previous = False
    for p in list(hdr.paragraphs):
        p._p.getparent().remove(p._p)
    tb = hdr.add_table(rows=1, cols=2, width=Twips(cw))
    _borders(tb, none=True)
    L, R = tb.rows[0].cells
    _set_widths(tb, [2300, cw - 2300])
    lp = L.paragraphs[0]
    if os.path.exists(LOGO):
        lp.add_run().add_picture(LOGO, width=Cm(2.6))
    rp = R.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_after = Pt(0)
    _font(rp.add_run(COMPANY), 9, bold=True, color=NAVY)
    rp2 = R.add_paragraph()
    rp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp2.paragraph_format.space_after = Pt(0)
    _font(rp2.add_run(f"{DEPARTMENT} · {title}"), 8, it=True, color=GREY)
    ep = hdr.add_paragraph()
    ep.paragraph_format.space_after = Pt(0)
    _pborder(ep, "bottom", BLUE, 8, 1)
    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]
    fp.text = ""
    fp.paragraph_format.tab_stops.add_tab_stop(Twips(cw), WD_TAB_ALIGNMENT.RIGHT)
    _font(fp.add_run("Tài liệu nội bộ · Phòng Kỹ thuật TEXO · Căn cứ NĐ 207/2026/NĐ-CP, Phụ lục VII"), 8, color=GREY)
    r = fp.add_run(); r._r.append(OxmlElement("w:tab"))
    _font(fp.add_run("Trang "), 9, color=GREY)
    _pagefield(fp)
    _pborder(fp, "top", "BFBFBF", 4, 4)


def _setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.left_margin = sec.right_margin = Twips(1134)
    sec.top_margin, sec.bottom_margin = Twips(1300), Twips(1000)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    _header_footer(sec, CW, "Kiểm tra hồ sơ chất lượng dự án")
    return sec


def _landscape(doc):
    sec = doc.add_section()
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(16838), Twips(11906)
    sec.left_margin = sec.right_margin = Twips(1134)
    sec.top_margin, sec.bottom_margin = Twips(1200), Twips(900)
    _header_footer(sec, CW_L, "Phụ lục danh mục đã đọc")
    return sec


# ----------------------------------------------------------------------------- nội dung
def xuat_docx(R, nguoi_kiem_tra="", ten_file="", ai=None):
    doc = Document()
    _setup(doc)
    now = datetime.now()
    meta = R["meta"]
    k = R["kpi"]

    para(doc, "", sa=2)
    para(doc, "PHIẾU KIỂM TRA HỒ SƠ CHẤT LƯỢNG DỰ ÁN", 17, True, NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
    para(doc, R["loai_nt_label"], 12, it=True, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
    sub = f"Ngày kiểm tra: {now:%d/%m/%Y} · Căn cứ: {R['nghi_dinh']} · File: {ten_file}"
    rp = para(doc, sub, 9.5, it=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, sa=8)
    _pborder(rp, "bottom", BLUE, 8, 4)

    # Kết luận
    sev = R["ket_luan_sev"]
    t = doc.add_table(rows=1, cols=1)
    _borders(t, color={"high": "C00000", "mid": "C45911", "ok": "2E7D32"}.get(sev, BLUE), sz=12)
    _cell_margins(t, 120, 120, 160, 160)
    c = t.rows[0].cells[0]
    _shade(c, FILL.get(sev, "EAF0F7"))
    p = c.paragraphs[0]
    _runs(p, f"{SEV_ICON.get(sev, '')} KẾT LUẬN: ", 12.5, True, FG.get(sev, NAVY))
    _runs(p, R["ket_luan"], 12.5, True, NAVY)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_before = Pt(4)
    _runs(p2, (f"Mục bắt buộc đã có trong danh mục: {k['muc_B_co']}/{k['muc_B']} ({k['ty_le_B']}%) · "
               f"Dòng hồ sơ chi tiết: {k['dong_chi_tiet']} · Có mã VPĐT: {k['ty_le_co_ma']}% · Số mã HS: {k['so_ma']}"
               + (f" (gom nhiều mục: {k['so_ma_gom']})" if k['so_ma_gom'] else "")
               + (f" · Khoảng thời gian hồ sơ: {k['tu_ngay']} – {k['den_ngay']}" if k['tu_ngay'] else "")), 10.5)
    _set_widths(t, [CW])
    para(doc, "", sa=4)

    # KPI
    kp = doc.add_table(rows=2, cols=5)
    _borders(kp, color="FFFFFF", sz=8)
    vals = [(f"{k['muc_B_co']}/{k['muc_B']}", "Mục bắt buộc có", "ok" if k['muc_B_co'] == k['muc_B'] else "mid"),
            (f"{k['ty_le_co_ma']}%", "Dòng có mã VPĐT", "ok" if k['ty_le_co_ma'] >= 90 else "mid"),
            (str(k["high"]), "🔴 Phải xử lý", "high" if k["high"] else "ok"),
            (str(k["mid"]), "🟠 Nên xử lý", "mid" if k["mid"] else "ok"),
            (str(k["info"]), "🔵 Lưu ý", "info")]
    for j, (v, lab, s) in enumerate(vals):
        a, b = kp.rows[0].cells[j], kp.rows[1].cells[j]
        for cc in (a, b):
            _shade(cc, FILL[s])
        pa = a.paragraphs[0]; pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(pa, v, 20, True, FG[s])
        pb = b.paragraphs[0]; pb.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(pb, lab, 9.5, True, NAVY)
    _set_widths(kp, [CW // 5] * 5)
    para(doc, "", sa=2)

    if R.get("warnings"):
        for w in R["warnings"]:
            para(doc, "⚠ " + w, 10.5, it=True, color=FG["mid"])

    # 1. Thông tin chung
    heading(doc, "1. THÔNG TIN CHUNG CỦA BỘ HỒ SƠ")
    lab = [("tieu_de", "Tên tài liệu"), ("du_an", "Dự án"), ("hang_muc", "Hạng mục"), ("goi_thau", "Gói thầu"),
           ("dia_diem", "Địa điểm"), ("chu_dau_tu", "Chủ đầu tư"), ("tvqlda", "Tư vấn QLDA"), ("tvgs", "Tư vấn giám sát"),
           ("tong_thau", "Tổng thầu"), ("nha_thau", "Nhà thầu thi công"), ("hop_dong", "Hợp đồng")]
    rows = [(l, meta.get(key) or "– chưa thể hiện –") for key, l in lab if meta.get(key) or key in ("du_an", "chu_dau_tu", "nha_thau")]
    rows.append(("Loại nghiệm thu", R["loai_nt_label"]))
    rows.append(("Vai trò nhà thầu", kb.VAI_TRO.get(R.get("vai_tro", "chinh"), "")))
    if R.get("ngay_nt"):
        rows.append(("Ngày nghiệm thu (dự kiến)", R["ngay_nt"].strftime("%d/%m/%Y")))
    kieu = {"DANH_MUC": "Danh mục hồ sơ", "BAO_CAO_DANG_TAI": "Báo cáo đăng tải hồ sơ VPĐT (nội bộ)"}.get(R["info"].get("kieu"), "Khác")
    rows.append(("Dạng tài liệu nhận diện", kieu))
    table(doc, ["Thông tin", "Nội dung"], rows, [2700, CW - 2700], size=10.5)

    # 2. Đối chiếu PL VII
    if R.get("tep"):
        heading(doc, "1b. CÁC DANH MỤC ĐÃ NẠP")
        para(doc, "Bộ hồ sơ được lập thành nhiều danh mục rời; mỗi danh mục được đọc riêng rồi gộp lại để đối chiếu.", 9.5, it=True, color=GREY)
        rows = [[i, t["ten"], t["loai_label"], t["dot"] or "", t["so_dong"],
                 ", ".join(f"{c} ({n})" for c, n in list((t.get("muc_dem") or {}).items())[:6]) or "—", t.get("so_ma", "")]
                for i, t in enumerate(R["tep"], 1)]
        table(doc, ["#", "Danh mục (file / sheet)", "Loại", "Đợt", "Số dòng", "Mục PL VII (số dòng)", "Mã HS"], rows,
              [400, 3000, 1700, 500, 700, 2638, 700], size=9)

    heading(doc, "2. ĐỐI CHIẾU VỚI DANH MỤC HỒ SƠ HOÀN THÀNH (PHỤ LỤC VII)")
    para(doc, "Yêu cầu: B – phải có hoặc phải nêu căn cứ · Đ – đối chiếu nếu có · Hình thức: VB – văn bản có số/ngày, "
              "PD – phải có văn bản phê duyệt/chấp thuận, Quyển – tập hồ sơ, danh mục chỉ là mục lục.", 9.5, it=True, color=GREY)
    hf = {"VB": "VB", "PD": "PD", "QUYEN": "Quyển"}
    for g in ("I", "II", "III", "T"):
        its = [it for it in R["items"] if it["nhom"] == g and (it["yeu_cau"] != "-" or it["so_dong"])]
        if not its:
            continue
        para(doc, f"{g}. {kb.NHOM[g]}", 11, True, NAVY, sa=2, sb=4).paragraph_format.keep_with_next = True
        rows = []
        for it in its:
            ma = ", ".join(it["ma"][:4]) + (f" (+{len(it['ma']) - 4})" if len(it["ma"]) > 4 else "")
            tim = f"{it['so_dong']} dòng" + (f" · {it['so_co_so_hieu']} có số/ngày" if it["so_dong"] else "")
            rows.append([it["code"], it["ten"] + ("\n(CĐT lưu)" if it.get("chu_the") == "CĐT" else ""), it["yeu_cau"],
                         hf[it["hinh_thuc"]], tim if it["so_dong"] else "—", ma or "—",
                         f"{SEV_ICON[it['trang_thai']]} {it['ket_luan']}", {"sev": it["trang_thai"], "cols": [6]}])
        table(doc, ["Mục", "Thành phần hồ sơ", "YC", "Hình thức", "Tìm thấy", "Mã VPĐT", "Kết luận"], rows,
              [700, 2900, 450, 800, 1250, 1388, 2150], size=9)

    # 3. Phát hiện
    heading(doc, "3. PHÁT HIỆN VÀ ĐỀ XUẤT")
    if R["findings"]:
        rows = [[f["stt"], f"{SEV_ICON[f['muc_do']]} {SEV_LABEL[f['muc_do']]}",
                 f["tieu_de"] + ("\n" + f["chi_tiet"] if f["chi_tiet"] else ""), f.get("de_xuat") or "",
                 {"sev": f["muc_do"], "cols": [1]}] for f in R["findings"]]
        table(doc, ["#", "Mức độ", "Phát hiện", "Đề xuất"], rows, [450, 1350, 4900, 2938], size=9.5, zebra=False)
    else:
        para(doc, "Không có phát hiện.", 11)

    # 4. Bản đồ mã VPĐT
    if R["ma_map"]:
        heading(doc, "4. BẢN ĐỒ MÃ HỒ SƠ TRÊN VPĐT")
        para(doc, "Mỗi mã HS trên VPĐT nên tương ứng một thành phần hồ sơ (một mục PL VII). Mã gom nhiều mục cho thấy hồ sơ "
                  "đang được nén chung, người tra cứu phải tải về toàn bộ mới tìm được tài liệu.", 10, it=True, color=GREY)
        rows = []
        for ma, e in R["ma_map"].items():
            codes = sorted(c for c in e["muc"] if c)
            nd = "; ".join(dict.fromkeys(e["noi_dung"]))[:180]
            s = "mid" if e.get("gom") else "ok"
            rows.append([ma, len(e["dong"]) + len(e["dong_ke_thua"]), ", ".join(codes) or "—", nd or "(mã của nhóm)",
                         ("🟠 Gom nhiều mục" if e.get("gom") else "🟢 Rõ"), {"sev": s, "cols": [4]}])
        table(doc, ["Mã HS", "Số dòng", "Mục PL VII", "Nội dung", "Đánh giá"], rows, [1500, 800, 2000, 3738, 1600], size=9)

    n = 5
    if ai and ai.get("ok"):
        d = ai["data"]
        heading(doc, f"{n}. Ý KIẾN BỔ SUNG (TRỢ LÝ AI)")
        n += 1
        para(doc, "Phần này do AI đọc kết quả đối chiếu của công cụ và đưa thêm nhận định; cần người kiểm tra thẩm định lại.",
             9.5, it=True, color=GREY)
        if d.get("nhan_xet_tong_quan"):
            para(doc, d["nhan_xet_tong_quan"], 11.5, sa=6)
        rr = d.get("rui_ro_bo_sung") or []
        if rr:
            rows = [[i, f"{SEV_ICON.get(x.get('muc_do'), '🔵')} {SEV_LABEL.get(x.get('muc_do'), 'Lưu ý')}", x.get("noi_dung", ""),
                     x.get("de_xuat", ""), {"sev": x.get("muc_do") if x.get("muc_do") in FILL else "info", "cols": [1]}]
                    for i, x in enumerate(rr, 1)]
            table(doc, ["#", "Mức độ", "Rủi ro", "Đề xuất"], rows, [450, 1350, 4400, 3438], size=9.5, zebra=False)
        if d.get("phan_loai_bo_sung"):
            para(doc, f"AI đã phân loại bổ sung {len(d['phan_loai_bo_sung'])} dòng (đánh dấu nguồn 'AI' trong phụ lục).", 10, it=True)

    if R.get("cau_hoi_muc"):
        heading(doc, f"{n}. CÂU HỎI KIỂM TRA THEO TỪNG MỤC")
        n += 1
        para(doc, "Các câu hỏi người kiểm tra cần đặt ra khi mở từng quyển hồ sơ – danh mục chỉ cho biết hồ sơ có hay chưa, "
                  "nội dung phải xác minh trên tài liệu.", 9.5, it=True, color=GREY)
        rows = []
        for q in R["cau_hoi_muc"]:
            code = q["code"] if q["code"] not in ("I", "III") else f"Phần {q['code']}"
            rows.append([code, q["ten"], "\n".join("– " + x for x in q["cau_hoi"]),
                         f"{SEV_ICON.get(q['trang_thai'], '')} {SEV_LABEL.get(q['trang_thai'], '')}" if q["code"] not in ("I", "III") else "",
                         {"sev": q["trang_thai"], "cols": [3]}])
        table(doc, ["Mục", "Thành phần", "Câu hỏi kiểm tra", "Hiện trạng"], rows, [800, 2300, 5138, 1400], size=9.5)

    heading(doc, f"{n}. ĐỀ NGHỊ ĐOÀN TƯ VẤN LÀM RÕ / BỔ SUNG")
    qs = list(R.get("cau_hoi", []))
    if ai and ai.get("ok"):
        for q in ai["data"].get("cau_hoi_cho_doan") or []:
            if q not in qs:
                qs.append(q)
    for i, q in enumerate(qs[:25], 1):
        para(doc, f"{i}. {q}", 11.5, sa=2)
    if not qs:
        para(doc, "Không có.", 11.5)

    # Ký
    para(doc, "", sa=6)
    para(doc, f"{CITY}, ngày {now:%d} tháng {now:%m} năm {now:%Y}", 11.5, it=True, align=WD_ALIGN_PARAGRAPH.RIGHT, sa=4)
    st = doc.add_table(rows=1, cols=2)
    _borders(st, none=True)
    for cell, title, name in ((st.rows[0].cells[0], "NGƯỜI KIỂM TRA", nguoi_kiem_tra),
                              (st.rows[0].cells[1], SIGNER_TITLE, SIGNER_NAME)):
        cell.text = ""
        p0 = cell.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(p0, title, 12, True)
        p1 = cell.add_paragraph(); p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(p1, "(Ký, ghi rõ họ tên)", 10.5, it=True)
        for _ in range(3):
            cell.add_paragraph()
        p2 = cell.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _runs(p2, name or "", 12.5, True)
    _set_widths(st, [CW // 2, CW // 2])
    for row in st.rows:
        _no_split(row)

    # Phụ lục – ngang
    _landscape(doc)
    para(doc, "PHỤ LỤC – DANH MỤC ĐÃ ĐỌC VÀ PHÂN LOẠI THEO PHỤ LỤC VII", 13, True, NAVY, sa=2)
    para(doc, "Nguồn phân loại: TK – theo từ khóa · KT – kế thừa nhóm cha · DM – theo loại danh mục (file) · MẪU – mục ghi sẵn trong mẫu · AI – trợ lý AI. Mã trong ngoặc là mã của nhóm cha (hồ sơ nằm chung trong mã đó).",
         9.5, it=True, color=GREY, sa=4)
    src = {"tu_khoa": "TK", "ke_thua": "KT", "ai": "AI", "tep": "DM", "mau": "MẪU"}
    rows = []
    for r in R["rows"]:
        ind = "   " * max(0, r["cap"] - 1)
        ma = ", ".join(r["ma"]) if r["ma"] else (("(" + ", ".join(r["ma_ke_thua"]) + ")") if r["ma_ke_thua"] else "")
        rows.append([r["id"], r["stt"], ind + r["noi_dung"][:260], r["so_vb"][:40], r["ngay_raw"], ma,
                     (r.get("muc") or "—") + (f" ({src.get(r.get('muc_nguon'), '')})" if r.get("muc") else "")])
    table(doc, ["ID", "STT", "Nội dung hồ sơ", "Số văn bản", "Ngày", "Mã VPĐT", "Mục PL VII"], rows,
          [550, 650, 6970, 2100, 1150, 1800, 1350], size=8.5)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
