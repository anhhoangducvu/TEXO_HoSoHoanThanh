# -*- coding: utf-8 -*-
"""
ĐỌC DANH MỤC HỒ SƠ từ Word (.docx), PDF có lớp chữ (không phải scan), Excel (.xlsx)
→ danh sách dòng chuẩn hóa + thông tin chung (dự án, hạng mục, CĐT, TVGS, nhà thầu, hợp đồng).

Không cần AI. Kết quả dùng cho core/checker.py.
"""
import io
import re
import unicodedata
from datetime import datetime, date

# ----------------------------------------------------------------------------- tiện ích
def bo_dau(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return s.replace("đ", "d").replace("Đ", "D")


def norm(s):
    s = bo_dau(s).lower()
    s = re.sub(r"[\r\n\t]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def clean(s, join=" "):
    if s is None:
        return ""
    if isinstance(s, (datetime, date)):
        return s.strftime("%d/%m/%Y")
    s = str(s).replace("\r", "\n")
    parts = [p.strip() for p in s.split("\n")]
    return re.sub(r"\s+", " ", join.join(p for p in parts if p)).strip()


# Mã hồ sơ / văn bản trên VPĐT (TEXO E-Office)
RE_MA = re.compile(r"(HS\s*[\.\-]?\s*\d{4}\s*[\.\-]\s*\d+|VB\s*(?:Đến|Đi|đến|đi|SP|KSSP|DEN|DI|den|di)?\s*[\-\.]?\s*\d{4}\s*[\-\.]\s*\d+)")
RE_NGAY = re.compile(r"\b(\d{1,2})\s*[/\.\-]\s*(\d{1,2})\s*[/\.\-]\s*(\d{4})\b")
RE_THANG = re.compile(r"^\s*(\d{1,2})\s*/\s*(\d{4})\s*$")
RE_STT = re.compile(r"^\s*([IVXLC]+|[A-ZĐ]|\d+(?:\.\d+)*|[a-zđ]\)|-|\+|•)\s*[\.\)]?\s*$")


def chuan_hoa_ma(raw):
    """'HS. 2025 . 3010' → 'HS.2025.3010'; 'VB Đến - 2026-89' → 'VB Đến-2026-89'."""
    s = re.sub(r"\s+", " ", raw.strip())
    m = re.match(r"HS\s*[\.\-]?\s*(\d{4})\s*[\.\-]\s*(\d+)", s, re.I)
    if m:
        return f"HS.{m.group(1)}.{m.group(2)}"
    m = re.match(r"VB\s*([^\d\-\.\s]*)\s*[\-\.]?\s*(\d{4})\s*[\-\.]\s*(\d+)", s, re.I)
    if m:
        loai = m.group(1).strip()
        return f"VB {loai}-{m.group(2)}-{m.group(3)}".replace("VB -", "VB-")
    return s


def tim_ma(text):
    out = []
    for m in RE_MA.finditer(str(text or "")):
        ma = chuan_hoa_ma(m.group(0))
        if ma not in out:
            out.append(ma)
    return out


def doc_ngay(v):
    """Trả (date|None, kiểu) — kiểu: 'dmy' | 'my' (chỉ tháng/năm) | 'us' (m/d/y nghi kiểu Mỹ) | 'loi' | ''."""
    if v is None or v == "":
        return None, ""
    if isinstance(v, datetime):
        return v.date(), "dmy"
    if isinstance(v, date):
        return v, "dmy"
    s = str(v).strip()
    m = RE_NGAY.search(s)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if mo > 12 and d <= 12:   # 3/23/2021 → kiểu Mỹ
            try:
                return date(y, d, mo), "us"
            except ValueError:
                return None, "loi"
        try:
            return date(y, mo, d), "dmy"
        except ValueError:
            return None, "loi"
    m = RE_THANG.match(s)
    if m:
        try:
            return date(int(m.group(2)), int(m.group(1)), 1), "my"
        except ValueError:
            return None, "loi"
    return None, ""


# ----------------------------------------------------------------------------- cột
COT_KHOA = {
    "stt": [r"^stt$", r"^tt$", r"^so tt$"],
    "noi_dung": [r"danh muc", r"noi dung", r"ten (ho so|tai lieu|van ban)", r"^ho so$", r"tai lieu"],
    "so_vb": [r"so (van ban|hieu|vb|ky hieu)", r"^so$", r"ky hieu"],
    "ngay": [r"ngay", r"thang"],
    "don_vi": [r"don vi", r"co quan ban hanh", r"ban hanh boi"],
    "ma": [r"texo", r"\bma\b", r"vpdt", r"e-office", r"ma so"],
    "tinh_trang": [r"tinh trang", r"ban (goc|sao)", r"hinh thuc"],
    "trang": [r"so trang", r"^trang$", r"so to"],
    "ghi_chu": [r"ghi chu"],
}


def nhan_dien_header(cells):
    """Trả dict cột→chỉ số nếu dòng là tiêu đề bảng; ngược lại None."""
    m = {}
    for j, c in enumerate(cells):
        t = norm(c)
        if not t or len(t) > 60:
            continue
        for k, pats in COT_KHOA.items():
            if k in m:
                continue
            if any(re.search(p, t) for p in pats):
                # 'ngay ban hanh' cũng chứa 'ban hanh' → ưu tiên ngày
                if k == "don_vi" and "ngay" in t:
                    continue
                if k == "noi_dung" and "ma" in t.split():
                    continue
                m[k] = j
                break
    if "noi_dung" in m and len(m) >= 2:
        return m
    return None


def suy_dien_cot(rows):
    """Bảng không có tiêu đề: đoán cột theo nội dung."""
    if not rows:
        return None
    n = max(len(r) for r in rows)
    score = {j: {"ma": 0, "ngay": 0, "stt": 0, "len": 0} for j in range(n)}
    for r in rows:
        for j, c in enumerate(r):
            t = clean(c)
            if not t:
                continue
            if tim_ma(t):
                score[j]["ma"] += 1
            if doc_ngay(t)[0]:
                score[j]["ngay"] += 1
            if RE_STT.match(t):
                score[j]["stt"] += 1
            score[j]["len"] += len(t)
    m = {}
    j_ma = max(score, key=lambda j: score[j]["ma"])
    if score[j_ma]["ma"]:
        m["ma"] = j_ma
    j_stt = max(score, key=lambda j: score[j]["stt"])
    if score[j_stt]["stt"] >= max(2, len(rows) // 4) and j_stt != m.get("ma"):
        m["stt"] = j_stt
    rest = [j for j in score if j not in m.values()]
    if rest:
        m["noi_dung"] = max(rest, key=lambda j: score[j]["len"])
    j_ng = max(score, key=lambda j: score[j]["ngay"])
    if score[j_ng]["ngay"] and j_ng not in m.values():
        m["ngay"] = j_ng
    return m if "noi_dung" in m else None


# ----------------------------------------------------------------------------- thông tin chung
META_KHOA = [
    ("du_an", r"^\s*(du an|ten du an|cong trinh)\b"),
    ("hang_muc", r"^\s*(hang muc|ten hang muc)\b"),
    ("goi_thau", r"^\s*(goi thau)\b"),
    ("dia_diem", r"^\s*(dia diem|dia diem xd|dia diem xay dung)\b"),
    ("chu_dau_tu", r"^\s*(dd |dai dien )?(chu dau tu)\b"),
    ("tvqlda", r"^\s*(tu van (qlda|quan ly du an))\b"),
    ("tvgs", r"^\s*(tu van giam sat)\b"),
    ("tong_thau", r"^\s*(tong thau)\b"),
    ("nha_thau", r"^\s*(nha thau thi cong|don vi thi cong|nha thau)\b"),
    ("hop_dong", r"^\s*(hop dong( so| thi cong xay dung so| xay dung so)?)\b"),
]


def doc_meta_dong(line, meta):
    t = norm(line)
    for k, p in META_KHOA:
        if k in meta:
            continue
        mm = re.match(p, t)
        if mm:
            # lấy phần sau dấu ':' trên bản có dấu
            raw = clean(line)
            if ":" in raw:
                val = raw.split(":", 1)[1].strip()
            else:
                val = raw[len(mm.group(0)):].strip(" :")
            val = val.strip(" :.")
            if val and len(val) > 2:
                meta[k] = val[:300]
            return True
    return False


def doc_meta_text(text, meta):
    lines = [l for l in str(text or "").split("\n") if l.strip()]
    for i, l in enumerate(lines[:80]):
        t = norm(l)
        if t.startswith("stt") or re.match(r"^(tt|so tt)\s", t):
            break  # hết phần thông tin chung, bắt đầu bảng
        if t.startswith("hop dong") and not re.search(r"\d+\s*/|\bso\b", t):
            continue
        doc_meta_dong(l, meta)
    # tiêu đề
    if "tieu_de" not in meta:
        for l in lines[:25]:
            t = norm(l)
            if ("danh muc" in t or "bao cao" in t or "ho so hoan thanh" in t) and len(t) < 160:
                meta["tieu_de"] = clean(l)
                break
    return meta


# ----------------------------------------------------------------------------- trích bảng theo định dạng
def _dedupe_merged(cells):
    """python-docx lặp lại ô gộp – bỏ ô trùng liền kề."""
    out = []
    prev = object()
    for c in cells:
        if c is not None and c == prev and str(c).strip():
            out.append("")
        else:
            out.append(c)
        prev = c
    return out


def tables_docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    tables = []
    for t in d.tables:
        rows = []
        for r in t.rows:
            cells = []
            seen = set()
            for c in r.cells:
                key = id(c._tc)
                if key in seen:
                    cells.append("")
                    continue
                seen.add(key)
                cells.append(c.text)
            rows.append(cells)
        tables.append(rows)
    paras = "\n".join(p.text for p in d.paragraphs)
    # header/footer section có thể chứa tên dự án
    return tables, paras, {"so_trang": None, "scan": False}


def tables_pdf(data):
    import pdfplumber
    tables, texts = [], []
    so_chu = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        n = len(pdf.pages)
        for p in pdf.pages:
            txt = p.extract_text() or ""
            texts.append(txt)
            so_chu += len(txt.strip())
            try:
                for t in p.extract_tables():
                    tables.append(t)
            except Exception:
                pass
    scan = so_chu < 40 * max(1, n)
    return tables, "\n".join(texts), {"so_trang": n, "scan": scan}


def tables_xlsx(data, sheets=None):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    tables, texts, names = [], [], []
    for ws in wb.worksheets:
        if sheets and ws.title not in sheets:
            continue
        rows = []
        for r in ws.iter_rows(values_only=True):
            if any(c not in (None, "") for c in r):
                rows.append(list(r))
        if rows:
            tables.append(rows)
            names.append(ws.title)
            texts.append("\n".join(" ".join(clean(c) for c in r if c not in (None, "")) for r in rows[:15]))
    return tables, "\n".join(texts), {"so_trang": None, "scan": False, "sheets": [w.title for w in wb.worksheets],
                                      "sheets_doc": names}


def danh_sach_sheet(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    return [w.title for w in wb.worksheets]


# ----------------------------------------------------------------------------- dòng
def _cap_do(stt):
    s = (stt or "").strip().rstrip(".")
    if not s:
        return None
    if re.fullmatch(r"[IVXLC]+", s):
        return 0
    if re.fullmatch(r"[A-ZĐ]", s):
        return 0
    if re.fullmatch(r"\d+(\.\d+)*", s):
        return 1 + s.count(".")
    if re.fullmatch(r"[a-zđ]\)?", s):
        return 3
    return None


def rows_from_tables(tables, meta, sheet_names=None):
    rows = []
    colmap = None
    for ti, t in enumerate(tables):
        t = [[("" if c is None else c) for c in r] for r in t]
        local_map = None
        start = 0
        for i, r in enumerate(t[:8]):
            h = nhan_dien_header([clean(c) for c in r])
            if h:
                local_map, start = h, i + 1
                # các dòng trước tiêu đề → thông tin chung
                for rr in t[:i]:
                    for c in rr:
                        for line in str(c or "").split("\n"):
                            doc_meta_dong(line, meta)
                    doc_meta_dong(" ".join(clean(c) for c in rr if clean(c)), meta)
                break
        if local_map is None:
            # bảng 2 cột 'Dự án | : ...' là thông tin chung
            if all(len([c for c in r if clean(c)]) <= 2 for r in t) and len(t) <= 12 and \
                    sum(doc_meta_dong(" ".join(clean(c) for c in r), meta) for r in t) >= max(1, len(t) // 2):
                continue
            # bảng tiếp nối từ trang trước, cùng số cột
            if colmap and max(len(r) for r in t) >= max(colmap.values()) + 1:
                local_map = colmap
            else:
                local_map = suy_dien_cot(t)
                if not local_map:
                    continue
        colmap = local_map
        sheet = sheet_names[ti] if sheet_names and ti < len(sheet_names) else ""
        for r in t[start:]:
            def g(k, join=" "):
                j = local_map.get(k)
                return clean(r[j], join) if j is not None and j < len(r) else ""
            if nhan_dien_header([clean(c) for c in r]):
                continue  # tiêu đề lặp lại ở trang sau
            nd = g("noi_dung")
            raw_all = " | ".join(clean(c) for c in r if clean(c))
            if not raw_all:
                continue
            # dòng ký tên
            if re.search(r"^(t/m|tm\.|dai dien|tu van giam sat|nha thau thi cong|chu dau tu)\b", norm(nd)) and \
                    not g("so_vb") and not g("ma") and len(nd) < 80 and not g("stt"):
                if re.search(r"t/m|gst|giam doc|\(ky", norm(raw_all)):
                    continue
            ma_cell = g("ma")
            mas = tim_ma(ma_cell) or ([] if "ma" in local_map else tim_ma(raw_all))
            # nếu cột 'ma' là cột ghi chú (mã nằm ở ghi chú)
            if not mas:
                mas = tim_ma(g("ghi_chu"))
            ngay_raw = r[local_map["ngay"]] if "ngay" in local_map and local_map["ngay"] < len(r) else ""
            rows.append(dict(
                stt=g("stt"), noi_dung=nd, so_vb=g("so_vb", ""), ngay_raw=clean(ngay_raw),
                ngay_obj=ngay_raw, don_vi=g("don_vi"), ma_raw=ma_cell, ma=mas,
                tinh_trang=g("tinh_trang"), trang=g("trang"), ghi_chu=g("ghi_chu"),
                nguon=(f"Sheet {sheet}" if sheet else f"Bảng {ti + 1}"), raw=raw_all))
    return rows, colmap


def rows_from_text(text):
    """Dự phòng khi PDF không tách được bảng: mỗi dòng có STT/mã/số hiệu là một mục."""
    rows = []
    for line in str(text or "").split("\n"):
        l = line.strip()
        if len(l) < 4:
            continue
        m = re.match(r"^([IVX]+|\d+(?:\.\d+)*|[A-Z])[\.\)]?\s+(.*)$", l)
        mas = tim_ma(l)
        if not m and not mas:
            if rows and not re.search(r"^(stt|danh muc|du an|dia diem)", norm(l)):
                rows[-1]["noi_dung"] += " " + l
                rows[-1]["raw"] += " " + l
            continue
        stt, rest = (m.group(1), m.group(2)) if m else ("", l)
        ng = RE_NGAY.search(rest)
        nd = RE_MA.sub("", rest)
        nd = RE_NGAY.sub("", nd).strip()
        rows.append(dict(stt=stt, noi_dung=nd, so_vb="", ngay_raw=ng.group(0) if ng else "", ngay_obj=ng.group(0) if ng else "",
                         don_vi="", ma_raw=" ".join(mas), ma=mas, tinh_trang="", trang="", ghi_chu="", nguon="Văn bản", raw=l))
    return rows


# ----------------------------------------------------------------------------- cấu trúc cây
def xay_cay(rows):
    """Gán cấp độ, cha, mã kế thừa từ dòng cha (nhóm hồ sơ gắn một mã)."""
    stack = []  # (level, idx)
    for i, r in enumerate(rows):
        r["id"] = i + 1
        lv = _cap_do(r["stt"])
        if lv is None:
            # dòng không số thứ tự: cùng cấp với dòng không-STT liền trước, hoặc con của dòng nhóm gần nhất
            if i > 0 and rows[i - 1].get("khong_stt") and stack:
                lv = rows[i - 1]["cap"]
            elif stack:
                par_lv, par_idx = stack[-1]
                par = rows[par_idx]
                par_group = not (par["so_vb"] or par["ngay_raw"])
                lv = par_lv + 1 if par_group else par_lv
            else:
                lv = 1
            r["khong_stt"] = True
        else:
            r["khong_stt"] = False
        while stack and stack[-1][0] >= lv:
            stack.pop()
        r["cap"] = lv
        r["cha"] = rows[stack[-1][1]]["id"] if stack else None
        r["nhom_lon"] = None
        for _, idx in stack:
            if rows[idx]["cap"] == 0:
                r["nhom_lon"] = rows[idx]["noi_dung"]
        # mã kế thừa
        r["ma_ke_thua"] = []
        if not r["ma"]:
            for _, idx in reversed(stack):
                if rows[idx]["ma"]:
                    r["ma_ke_thua"] = rows[idx]["ma"]
                    break
        stack.append((lv, i))
    ids = {r["id"]: r for r in rows}
    for r in rows:
        r["so_con"] = 0
    for r in rows:
        if r["cha"]:
            ids[r["cha"]]["so_con"] += 1
    return rows


# ----------------------------------------------------------------------------- API
def doc_file(name, data, sheets=None):
    """Đọc một file → dict(meta, rows, warnings, text, info)."""
    ext = name.lower().rsplit(".", 1)[-1]
    warnings = []
    meta = {}
    sheet_names = None
    if ext == "docx":
        tables, text, info = tables_docx(data)
    elif ext == "pdf":
        tables, text, info = tables_pdf(data)
        if info["scan"]:
            warnings.append("PDF dạng ảnh scan (gần như không có lớp chữ) – công cụ chưa đọc được. "
                            "Đề nghị dùng bản Word/Excel gốc hoặc PDF xuất trực tiếp từ Word/Excel (AI cũng không đọc được ảnh scan).")
    elif ext in ("xlsx", "xlsm"):
        tables, text, info = tables_xlsx(data, sheets)
        sheet_names = info.get("sheets_doc")
    elif ext == "doc":
        return dict(meta={}, rows=[], warnings=["File .doc (Word 97-2003) – vui lòng mở bằng Word và lưu lại thành .docx."],
                    text="", info={})
    else:
        return dict(meta={}, rows=[], warnings=[f"Định dạng .{ext} chưa hỗ trợ (dùng .docx, .pdf, .xlsx)."], text="", info={})

    doc_meta_text(text, meta)
    rows, colmap = rows_from_tables(tables, meta, sheet_names)
    if len(rows) < 3 and text and not info.get("scan"):
        rows_t = rows_from_text(text)
        if len(rows_t) > len(rows):
            rows = rows_t
            warnings.append("Không nhận được bảng – đã đọc theo từng dòng chữ (độ chính xác thấp hơn, nên kiểm tra lại).")
    # lọc dòng rác: không có nội dung và không có mã
    rows = [r for r in rows if (r["noi_dung"] or r["ma"] or r["so_vb"])]
    rows = xay_cay(rows)
    info["co_cot"] = sorted(colmap.keys()) if colmap else []
    t = norm(meta.get("tieu_de", "") + " " + text[:600])
    if "dang tai" in t or "e-office" in t or "texo e" in t:
        info["kieu"] = "BAO_CAO_DANG_TAI"
    elif "danh muc" in t:
        info["kieu"] = "DANH_MUC"
    else:
        info["kieu"] = "KHAC"
    return dict(meta=meta, rows=rows, warnings=warnings, text=text, info=info)


def doan_loai_nt(meta, text):
    """Đoán loại nghiệm thu từ tiêu đề: HT / GT / GD / None."""
    t = norm((meta.get("tieu_de") or "") + " " + str(text or "")[:400])
    if re.search(r"hoan thanh (hang muc|cong trinh)|dua vao su dung|quyet toan", t):
        return "HT"
    if re.search(r"hoan thanh goi thau", t):
        return "GT"
    if re.search(r"giai doan|bo phan", t):
        return "GD"
    return None


def doc_danh_sach_vpdt(name, data):
    """File xuất từ VPĐT (Danh-sach-ho-so.xlsx …): trả {mã: mô tả} để kiểm tra mã có tồn tại."""
    ext = name.lower().rsplit(".", 1)[-1]
    out = {}
    if ext in ("xlsx", "xlsm"):
        tables, _, _ = tables_xlsx(data)
    elif ext == "docx":
        tables, _, _ = tables_docx(data)
    elif ext == "pdf":
        tables, _, _ = tables_pdf(data)
    else:
        return out
    for t in tables:
        for r in t:
            cells = [clean(c) for c in r]
            line = " | ".join(c for c in cells if c)
            for ma in tim_ma(line):
                out.setdefault(ma, line[:250])
    return out


def doc_tu_ai(data, warnings=None):
    """Chuyển JSON AI trích xuất {meta, rows[]} → cấu trúc như doc_file()."""
    meta = {k: v for k, v in (data.get("meta") or {}).items() if v}
    rows = []
    for x in data.get("rows") or []:
        nd = clean(x.get("noi_dung"))
        ma_raw = clean(x.get("ma"))
        rows.append(dict(stt=clean(x.get("stt")), noi_dung=nd, so_vb=clean(x.get("so_vb"), ""),
                         ngay_raw=clean(x.get("ngay")), ngay_obj=clean(x.get("ngay")), don_vi=clean(x.get("don_vi")),
                         ma_raw=ma_raw, ma=tim_ma(ma_raw), tinh_trang=clean(x.get("tinh_trang")), trang="",
                         ghi_chu=clean(x.get("ghi_chu")), nguon="AI trích xuất", raw=nd))
    rows = xay_cay([r for r in rows if r["noi_dung"] or r["ma"]])
    co = ["noi_dung", "stt", "so_vb", "ngay", "don_vi"] + (["ma"] if any(r["ma"] for r in rows) else []) + \
         (["tinh_trang"] if any(r["tinh_trang"] for r in rows) else [])
    t = norm(meta.get("tieu_de", ""))
    kieu = "BAO_CAO_DANG_TAI" if ("dang tai" in t or "e-office" in t) else "DANH_MUC"
    return dict(meta=meta, rows=rows, warnings=(warnings or []) + ["Danh mục do AI trích xuất từ văn bản – cần đối chiếu lại bản gốc."],
                text="", info={"co_cot": co, "kieu": kieu})
